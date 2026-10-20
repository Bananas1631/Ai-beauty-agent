"""Месячная выгрузка для счёта по тарифу «Результат» в Excel: pandas собирает, openpyxl оформляет.

Листы:
- «Сводка» — записи через ИИ, тариф, сумма к оплате, расход на ИИ. Итоги — формулами Excel над
  листом «Записи»: бухгалтер салона видит, из чего сложилась сумма, и может проверить сам;
- «Записи» — одна строка на оплачиваемую запись (правило domain/billing.py) и строка «Итого»;
- «По услугам» — сводная таблица pandas: сколько записей и на какую сумму по каждой услуге.

Деньги — Decimal от базы до ячейки: pandas держит их в столбце object (без float), в ячейку уходит
Decimal, формат «# ##0,00 ₽». Время визита — по часовому поясу салона, без пояса в ячейке:
Excel не умеет даты с поясом. Для выгрузок больше 50 тысяч строк — write_only (поток без
стилей по ячейкам), здесь это не нужно: у салона сотни записей в месяц.
"""

import io
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from beauty_agent.analytics.queries import Analytics
from beauty_agent.db.models import Salon
from beauty_agent.domain.billing import PRICE_PER_BOOKING_RUB, BillingPeriod
from beauty_agent.domain.errors import DomainValidationError
from beauty_agent.reports.invoice import InvoiceLine, billable_bookings, llm_cost

MONEY = '#,##0.00 "₽"'
DATETIME = "DD.MM.YYYY HH:MM"
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
STATUS_RU = {"confirmed": "подтверждена", "completed": "визит состоялся", "no_show": "клиент не пришёл"}
RECORDS = "Записи"
RULE = "Правило: запись создана ИИ, визит в этом месяце наступил и не отменён (договор, «Порядок расчёта»)."


@dataclass(frozen=True, slots=True)
class InvoiceData:
    salon_name: str
    month: str  # ГГГГ-ММ
    timezone: str
    lines: Sequence[InvoiceLine]
    llm_cost_rub: Decimal
    price_rub: Decimal = PRICE_PER_BOOKING_RUB


_MONTH = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def parse_month(month: str) -> tuple[int, int]:
    match = _MONTH.fullmatch(month)
    if match is None:
        raise DomainValidationError("Месяц в формате ГГГГ-ММ, например 2026-11")
    return int(match[1]), int(match[2])


def invoice_filename(salon: Salon, month: str) -> str:
    return f"{salon.slug}-{month}-schet.xlsx"


async def collect_invoice(
    session: AsyncSession, salon: Salon, month: str, *, now: datetime, analytics: Analytics | None = None
) -> InvoiceData:
    """Записи и итог — из PostgreSQL (источник правды для денег), расход на ИИ — из ClickHouse, если он есть."""
    year, number = parse_month(month)
    period = BillingPeriod.month(year, number, salon.timezone)
    lines = await billable_bookings(session, salon.id, period, now)
    cost = (
        await analytics.salon_llm_cost(salon.id, period.start, period.end)
        if analytics is not None
        else await llm_cost(session, salon.id, period.start, period.end)
    )
    return InvoiceData(salon.name, month, salon.timezone, lines, cost)


def records_frame(data: InvoiceData) -> pd.DataFrame:
    tz = ZoneInfo(data.timezone)
    return pd.DataFrame(
        {
            "Дата визита": [line.starts_at.astimezone(tz).replace(tzinfo=None) for line in data.lines],
            "Услуга": [line.service for line in data.lines],
            "Мастер": [line.master for line in data.lines],
            "Статус": [STATUS_RU.get(line.status.value, line.status.value) for line in data.lines],
            "ID записи в CRM": [line.external_id or "" for line in data.lines],
            "Сумма, ₽": pd.Series([data.price_rub] * len(data.lines), dtype=object),
        }
    )


def by_service(records: pd.DataFrame) -> pd.DataFrame:
    """Сводная по услугам: число записей и сумма. Сумма Decimal — через Python sum, не float."""
    if records.empty:
        return pd.DataFrame({"Услуга": [], "Записей": [], "Сумма, ₽": []})
    grouped = records.groupby("Услуга", sort=True).agg(
        Записей=("Сумма, ₽", "size"), **{"Сумма, ₽": ("Сумма, ₽", lambda values: sum(values, Decimal(0)))}
    )
    return grouped.reset_index().sort_values(["Записей", "Услуга"], ascending=[False, True], ignore_index=True)


def _header(sheet: Worksheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(sheet: Worksheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_frame(sheet: Worksheet, frame: pd.DataFrame) -> None:
    sheet.append(list(frame.columns))
    for values in frame.itertuples(index=False):
        sheet.append(list(values))


def invoice_workbook(data: InvoiceData) -> bytes:
    records = records_frame(data)
    count = len(records)
    first, last = 2, count + 1  # строки данных на листе «Записи»
    book = Workbook()

    summary = book.active
    assert summary is not None  # noqa: S101 — у новой книги активный лист есть всегда
    summary.title = "Сводка"
    amount_range = f"'{RECORDS}'!F{first}:F{max(last, first)}"
    rows: list[tuple[str, object]] = [
        ("Салон", data.salon_name),
        ("Месяц визитов", data.month),
        ("Записей через ИИ", f"=COUNTA('{RECORDS}'!A{first}:A{max(last, first)})" if count else 0),
        ("Тариф за запись, ₽", data.price_rub),
        ("Сумма к оплате, ₽", f"=SUM({amount_range})" if count else Decimal(0)),
        ("Проверка: записи × тариф, ₽", "=B4*B5"),
        ("Расход на ИИ за месяц, ₽", data.llm_cost_rub),
    ]
    summary.append(["Показатель", "Значение"])
    for label, value in rows:
        summary.append([label, value])
    _header(summary)
    for cell in ("B5", "B6", "B7", "B8"):
        summary[cell].number_format = MONEY
    summary["A10"] = RULE
    summary["A10"].font = Font(italic=True, color="666666")
    _widths(summary, (34, 28))

    sheet = book.create_sheet(RECORDS)
    _write_frame(sheet, records)
    _header(sheet)
    for row in sheet.iter_rows(min_row=first, max_row=last):
        row[0].number_format = DATETIME
        row[5].number_format = MONEY
    total_row = last + 1
    sheet.cell(total_row, 1, "Итого")
    sheet.cell(total_row, 5, f"=COUNTA(A{first}:A{last})" if count else 0)
    total = sheet.cell(total_row, 6, f"=SUM(F{first}:F{last})" if count else Decimal(0))
    total.number_format = MONEY
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"  # заголовок остаётся на месте при прокрутке
    sheet.auto_filter.ref = f"A1:F{max(last, 1)}"
    _widths(sheet, (18, 36, 22, 18, 18, 14))

    services = book.create_sheet("По услугам")
    _write_frame(services, by_service(records))
    _header(services)
    for row in services.iter_rows(min_row=2):
        row[2].number_format = MONEY
    services.freeze_panes = "A2"
    _widths(services, (36, 12, 14))

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
