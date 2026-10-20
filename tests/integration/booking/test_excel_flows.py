"""Excel в работе: прайс салона в каталог, выгрузка счёта владельцу по API и рассылка 1-го числа."""

import io
import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal

from fastapi import FastAPI
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker
from tests.integration.api.test_admin_api import bearer, bound_admin
from tests.integration.booking.test_reports import MSK, add_booking, at, catalog_salon  # noqa: F401 — фикстура

from beauty_agent.catalog.import_excel import import_price_list, parse_price_list
from beauty_agent.db.models import Salon, Service
from beauty_agent.db.repositories.admins import AdminsRepository
from beauty_agent.domain.enums import AdminRole, BookingStatus
from beauty_agent.workers.reports import previous_month, send_monthly_invoices

CSV = "Услуга;Цена;Длительность\nМаникюр;1 500;60\nПедикюр;2000–2500;1:30\n"


async def services_of(session: AsyncSession, salon: Salon) -> dict[str, Service]:
    rows = await session.scalars(select(Service).where(Service.salon_id == salon.id))
    return {service.name: service for service in rows}


async def test_price_list_creates_updates_and_deactivates(session: AsyncSession, salon: Salon) -> None:
    first = await import_price_list(session, salon.id, parse_price_list(CSV.encode("cp1251"), "price.csv"))
    await session.commit()
    updated_csv = "Услуга;Цена;Длительность\nМаникюр;1 700;60\n"
    second = await import_price_list(session, salon.id, parse_price_list(updated_csv.encode(), "price.csv"))
    await session.commit()

    services = await services_of(session, salon)
    assert (first.created, first.updated) == (2, 0)
    assert (second.created, second.updated, second.deactivated) == (0, 1, 1)
    assert services["Маникюр"].price_min == services["Маникюр"].price_max == 1700
    assert services["Педикюр"].is_active is False  # на услугу могли ссылаться записи — не удаляем
    assert (services["Педикюр"].price_min, services["Педикюр"].price_max) == (Decimal(2000), Decimal(2500))


async def test_price_list_with_errors_is_not_imported(session: AsyncSession, salon: Salon) -> None:
    broken = parse_price_list("Услуга;Цена;Длительность\nМаникюр;;60\n".encode(), "price.csv")

    try:
        await import_price_list(session, salon.id, broken)
    except ValueError:
        pass
    else:  # pragma: no cover
        raise AssertionError("прайс с ошибками импортирован")
    assert await services_of(session, salon) == {}


async def test_owner_downloads_invoice_and_admin_does_not(
    app: FastAPI,
    client: AsyncClient,
    session: AsyncSession,
    catalog_salon: Salon,  # noqa: F811
) -> None:
    # прошлый месяц относительно настоящих часов: оплачивается только наступивший визит
    month_start = datetime.now(MSK).replace(day=1, hour=12, minute=0, second=0, microsecond=0)
    visit = (month_start - timedelta(days=20)).replace(day=10)
    month = f"{visit:%Y-%m}"
    await add_booking(session, catalog_salon, starts_at=visit, status=BookingStatus.COMPLETED)
    await add_booking(session, catalog_salon, starts_at=visit + timedelta(days=1), status=BookingStatus.CANCELLED)
    await bound_admin(session, catalog_salon)  # владелец
    await AdminsRepository(session).bind(catalog_salon.id, "777", "777", "Ирина", AdminRole.ADMIN)
    await session.commit()

    owner = await client.get("/api/v1/admin/invoice.xlsx", params={"month": month}, headers=bearer(app, catalog_salon))
    admin = await client.get(
        "/api/v1/admin/invoice.xlsx", params={"month": month}, headers=bearer(app, catalog_salon, user_id="777")
    )
    bad_month = await client.get(
        "/api/v1/admin/invoice.xlsx", params={"month": "2026-13"}, headers=bearer(app, catalog_salon)
    )

    assert owner.status_code == 200
    assert owner.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f'filename="{catalog_salon.slug}-{month}-schet.xlsx"' in owner.headers["content-disposition"]
    records = load_workbook(io.BytesIO(owner.content))["Записи"]
    assert records.max_row == 3  # заголовок, одна оплачиваемая запись, «Итого»
    assert admin.status_code == 403
    assert bad_month.status_code == 422


async def test_monthly_invoices_go_to_each_salon(
    session: AsyncSession,
    session_factory: async_sessionmaker[AsyncSession],
    catalog_salon: Salon,  # noqa: F811
) -> None:
    await add_booking(session, catalog_salon, starts_at=at(10, 12), status=BookingStatus.COMPLETED)
    sent: list[tuple[uuid.UUID, str, bytes, str]] = []

    async def send(salon_id: uuid.UUID, name: str, content: bytes, caption: str) -> int:
        sent.append((salon_id, name, content, caption))
        return 1

    count = await send_monthly_invoices(session_factory, send, now=datetime(2026, 12, 1, 6, tzinfo=UTC))

    ours = [item for item in sent if item[0] == catalog_salon.id]
    assert count >= 1
    assert len(ours) == 1
    assert ours[0][1] == f"{catalog_salon.slug}-2026-11-schet.xlsx"
    assert "записей через ИИ: 1" in ours[0][3]


def test_previous_month() -> None:
    assert previous_month(datetime(2027, 1, 1, tzinfo=UTC)) == "2026-12"
    assert previous_month(datetime(2026, 3, 31, tzinfo=UTC)) == "2026-02"
