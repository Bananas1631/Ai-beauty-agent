"""Excel: итоги формулами совпадают с данными, деньги без float; импорт прайса с ошибками по строкам."""

import io
from datetime import UTC, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from beauty_agent.catalog.import_excel import (
    PriceItem,
    normalize_header,
    parse_duration,
    parse_money,
    parse_price_list,
)
from beauty_agent.domain.enums import BookingStatus
from beauty_agent.domain.errors import DomainValidationError
from beauty_agent.reports.excel import InvoiceData, by_service, invoice_workbook, parse_month, records_frame
from beauty_agent.reports.invoice import InvoiceLine

START = datetime(2026, 11, 3, 7, 0, tzinfo=UTC)


def lines(count: int) -> list[InvoiceLine]:
    services = ["Маникюр с покрытием", "Стрижка женская", "Маникюр с покрытием"]
    return [
        InvoiceLine(
            starts_at=START + timedelta(days=index),
            service=services[index % 3],
            master="Анна",
            status=BookingStatus.COMPLETED if index % 4 else BookingStatus.NO_SHOW,
            external_id=str(1000 + index),
        )
        for index in range(count)
    ]


def invoice(count: int) -> InvoiceData:
    return InvoiceData("Студия Лаванда", "2026-11", "Europe/Moscow", lines(count), Decimal("123.4567"))


def test_totals_in_formulas_match_data() -> None:
    book = load_workbook(io.BytesIO(invoice_workbook(invoice(7))))
    summary, records = book["Сводка"], book["Записи"]

    assert summary["B4"].value == "=COUNTA('Записи'!A2:A8)"
    assert summary["B6"].value == "=SUM('Записи'!F2:F8)"
    assert summary["B7"].value == "=B4*B5"
    assert summary["B5"].value == 59
    assert summary["B8"].value == pytest.approx(123.4567)
    # формулы ссылаются ровно на строки данных: посчитаем то, что посчитает Excel
    amounts = [records.cell(row, 6).value for row in range(2, 9)]
    assert sum(Decimal(str(value)) for value in amounts) == Decimal(59 * 7)
    assert records["A9"].value == "Итого"
    assert records["F9"].value == "=SUM(F2:F8)"
    assert records.freeze_panes == "A2"
    assert records["F2"].number_format == '#,##0.00 "₽"'


def test_visit_time_is_local_and_naive() -> None:
    records = load_workbook(io.BytesIO(invoice_workbook(invoice(1))))["Записи"]

    assert records["A2"].value == datetime(2026, 11, 3, 10, 0)  # 07:00 UTC = 10:00 по Москве
    assert records["D2"].value == "клиент не пришёл"


def test_by_service_counts_and_sums_with_decimal() -> None:
    frame = by_service(records_frame(invoice(5)))

    assert list(frame["Услуга"]) == ["Маникюр с покрытием", "Стрижка женская"]
    assert list(frame["Записей"]) == [3, 2]
    assert list(frame["Сумма, ₽"]) == [Decimal(177), Decimal(118)]
    assert all(isinstance(value, Decimal) for value in frame["Сумма, ₽"])


def test_empty_month_has_zeros_not_broken_formulas() -> None:
    book = load_workbook(io.BytesIO(invoice_workbook(invoice(0))))

    assert book["Сводка"]["B4"].value == 0
    assert book["Сводка"]["B6"].value == 0
    assert book["Записи"]["A2"].value == "Итого"
    assert book["По услугам"].max_row == 1


@pytest.mark.parametrize("month", ["2026-13", "26-11", "2026-1", ""])
def test_bad_month(month: str) -> None:
    with pytest.raises(DomainValidationError):
        parse_month(month)


# ---------------------------------------------------------------- импорт прайса


def xlsx(rows: list[list[object]]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("1500", (Decimal("1500.00"), Decimal("1500.00"))),
        ("1 500 ₽", (Decimal("1500.00"), Decimal("1500.00"))),
        ("1 500,50 руб.", (Decimal("1500.50"), Decimal("1500.50"))),
        ("от 2000", (Decimal("2000.00"), Decimal("2000.00"))),
        ("2000–2500", (Decimal("2000.00"), Decimal("2500.00"))),
    ],
)
def test_money(raw: str, expected: tuple[Decimal, Decimal]) -> None:
    assert parse_money(raw) == expected


@pytest.mark.parametrize("raw", ["бесплатно", "0", "-100", "2500–2000", "10.555", "5000000"])
def test_bad_money(raw: str) -> None:
    with pytest.raises(ValueError, match="цена"):
        parse_money(raw)


@pytest.mark.parametrize(
    ("raw", "minutes"), [("90", 90), ("1:30", 90), ("1 ч 30 мин", 90), ("2 часа", 120), ("45 мин.", 45), ("1,5 ч", 90)]
)
def test_duration(raw: str, minutes: int) -> None:
    assert parse_duration(raw) == minutes


@pytest.mark.parametrize("raw", ["полтора часа", "0", "800", "3"])
def test_bad_duration(raw: str) -> None:
    with pytest.raises(ValueError, match="длительность"):
        parse_duration(raw)


def test_headers_are_matched_loosely() -> None:
    assert normalize_header("  Цена, ₽ ") == "цена"
    assert normalize_header("Длительность (мин.)") == "длительность мин"


def test_xlsx_price_list() -> None:
    data = xlsx(
        [
            ["Название услуги", "Категория", "Цена, ₽", "Длительность, мин"],
            ["Маникюр с покрытием", "Ногти", "2 500", "90"],
            [None, None, None, None],  # пустая строка-разделитель
            ["Стрижка женская", "Волосы", "от 1800", "1:00"],
        ]
    )

    result = parse_price_list(data, "прайс.xlsx")

    assert result.ok
    assert result.items == [
        PriceItem(2, "Маникюр с покрытием", Decimal("2500.00"), Decimal("2500.00"), 90, "Ногти"),
        PriceItem(4, "Стрижка женская", Decimal("1800.00"), Decimal("1800.00"), 60, "Волосы"),
    ]
    assert result.items[0].external_id == PriceItem(9, "маникюр  с покрытием", Decimal(1), Decimal(1), 5).external_id


def test_cp1251_csv_with_semicolons() -> None:
    text = "Услуга;Цена;Время\r\nБрови: коррекция;800;30\r\nОкрашивание бровей;1 200 руб.;45\r\n"

    result = parse_price_list(text.encode("cp1251"), "price.csv")

    assert result.ok
    assert [(item.name, item.price_min, item.duration_min) for item in result.items] == [
        ("Брови: коррекция", Decimal("800.00"), 30),
        ("Окрашивание бровей", Decimal("1200.00"), 45),
    ]


def test_all_row_errors_are_reported_with_excel_row_numbers() -> None:
    data = xlsx(
        [
            ["Услуга", "Цена", "Длительность"],
            ["Маникюр", "1500", "60"],
            ["", "900", "30"],
            ["Педикюр", "бесплатно", "полтора часа"],
            ["Маникюр", "1600", "60"],
        ]
    )

    result = parse_price_list(data, "price.xlsx")

    assert not result.ok
    assert [str(error) for error in result.errors] == [
        "строка 3: пустое название услуги",
        "строка 4: цена «бесплатно» — не число",
        "строка 4: длительность «полтора часа» — укажите минуты: 90, 1:30 или «1 ч 30 мин»",
        "строка 5: услуга «Маникюр» уже есть в строке 2",
    ]


def test_missing_columns_and_unknown_format() -> None:
    no_duration = parse_price_list(xlsx([["Услуга", "Цена"], ["Маникюр", "1500"]]), "price.xlsx")
    pdf = parse_price_list(b"%PDF-1.7", "price.pdf")
    broken = parse_price_list(b"not a zip", "price.xlsx")

    assert [str(error) for error in no_duration.errors] == ["строка 1: нет колонок: Длительность"]
    assert "xlsx" in str(pdf.errors[0])
    assert str(broken.errors[0]).startswith("Файл не прочитан")
